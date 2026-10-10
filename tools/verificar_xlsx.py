"""Compara el .xlsx recalculado por LibreOffice con los valores del motor JS."""
import json, sys
import openpyxl
plan = json.load(open(sys.argv[1]))
wb = openpyxl.load_workbook(sys.argv[2], data_only=True)
errores, malos = [], 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and (v.startswith('#') or v.startswith('Err:')):
                errores.append(f"{ws.title}!{c.coordinate}={v}")
print('celdas con error:', len(errores), errores[:15])
fl = wb['Flujo']; m = plan['motor']; n = m['n']
def cmp(nombre, a, b, tol):
    global malos
    ok = (a is None and b in (None, 'No existe', 'No recupera')) or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) <= tol)
    if not ok: malos += 1
    print(('OK ' if ok else 'MAL'), nombre, a, b)
for alt, b in plan['filas'].items():
    if not m.get(alt): continue
    fc = b['rows']['fc']
    for t in range(n + 1):
        cmp(f'{alt} flujo año {t}', m[alt]['flujo'][t], fl.cell(fc, 2 + t).value, 0.01)
    for k in ['van', 'tir', 'tur', 'bc', 'tri', 'vac', 'cae']:
        cmp(f'{alt} {k}', m[alt][k], fl[f"Y{b['ind'][k]}"].value, 0.01 if k in ('van','vac','cae') else 1e-6)
    cmp(f'{alt} precio mínimo', m[alt]['precioMin'][0], fl.cell(12, 26 + list(plan['filas']).index(alt)).value, 1e-6)
if plan['preguntas']:
    rg = wb['Riesgo']; q = plan['riesgo']['preg']
    for i, p in enumerate(plan['preguntas']):
        cmp(f'pregunta {i} empírica', p['emp'], rg[f'D{q+i}'].value, 1e-9)
        cmp(f'pregunta {i} normal', p['nor'], rg[f'E{q+i}'].value, 1e-6)
print('DIFERENCIAS:', malos, '· ERRORES:', len(errores))
