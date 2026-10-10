"""Convierte el plan JSON del complemento en un .xlsx (fórmulas + formato básico).
Sirve para recalcular con LibreOffice y verificar las fórmulas sin Excel.
Uso: python3 plan_a_xlsx.py plan.json salida.xlsx"""
import json, re, sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import range_boundaries, get_column_letter

XLFN = ['PERCENTILE.INC', 'NORM.DIST', 'NORM.INV', 'STDEV.S']

def xlfn(f):
    for fn in XLFN:
        f = re.sub(r'(?<![\w.])' + re.escape(fn) + r'\(', '_xlfn.' + fn + '(', f)
    return f

def celdas(ws, addr):
    c1, r1, c2, r2 = range_boundaries(addr if ':' in addr else addr + ':' + addr)
    return [[ws.cell(r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]

plan = json.load(open(sys.argv[1]))
wb = openpyxl.Workbook()
wb.remove(wb.active)
for h in plan['hojas']:
    ws = wb.create_sheet(h['nombre'])
    if h.get('oculta'): ws.sheet_state = 'hidden'
    for op in h['ops']:
        t = op['op']
        if t == 'set':
            filas = celdas(ws, op['addr'])
            for fila, datos in zip(filas, op['datos']):
                for c, v in zip(fila, datos):
                    if v is None or v == '':
                        c.value = None
                    elif isinstance(v, str) and v.startswith('='):
                        c.value = xlfn(v)
                    else:
                        c.value = v
        elif t == 'fmt':
            f = op['f']
            for fila in celdas(ws, op['addr']):
                for c in fila:
                    if 'fill' in f: c.fill = PatternFill('solid', fgColor=f['fill'].lstrip('#'))
                    if f.get('bold') or f.get('size') or f.get('italic'):
                        c.font = Font(bold=f.get('bold', False), size=f.get('size', 11), italic=f.get('italic', False))
                    if 'numFmt' in f: c.number_format = f['numFmt']
                    if f.get('wrap'): c.alignment = Alignment(wrap_text=True)
        elif t == 'nombre':
            wb.defined_names[op['nombre']] = DefinedName(op['nombre'], attr_text=op['ref'])
# Columnas y filas ocultas (lo mismo que hace ajustarVista en Excel)
for nombre, cfg in (plan.get('vista') or {}).items():
    if not isinstance(cfg, dict) or nombre not in wb.sheetnames: continue
    ws = wb[nombre]
    if cfg.get('columnas'):
        a, b = cfg['columnas'].split(':')
        for c in range(range_boundaries(a + '1')[0], range_boundaries(b + '1')[0] + 1):
            ws.column_dimensions[get_column_letter(c)].hidden = True
    if cfg.get('filas'):
        a, b = map(int, cfg['filas'].split(':'))
        for r in range(a, b + 1):
            ws.row_dimensions[r].hidden = True
wb.save(sys.argv[2])
print('xlsx ok')
